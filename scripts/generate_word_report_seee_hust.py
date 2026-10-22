"""Tự động xuất file Word (.docx) Báo cáo Đồ án Chương 1, 2, 3 theo đúng chuẩn format
Trường Điện - Điện tử (SEEE) — Đại học Bách khoa Hà Nội (HUST):
- Khổ giấy A4 (210 x 297 mm)
- Lề: Trái 3.0 cm, Phải 2.0 cm, Trên 2.0 cm, Dưới 2.0 cm
- Font chữ: Times New Roman, 13 pt, giãn dòng 1.3 lines, căn đều 2 bên (Justify)
- Tên Chương: 16 pt, IN HOA ĐẬM, ngắt trang mới
- Mục cấp 1 (1.1, 2.1...): 14 pt, In đậm
- Mục cấp 2 (1.1.1, 2.1.1...): 13 pt, In đậm
- Bảng biểu và Hình ảnh thực tế chèn trực tiếp đúng quy chuẩn Bách Khoa.
"""

from __future__ import annotations

import pathlib
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "kenh_do_an_thay_duyet"
OUT_DOCX = OUT_DIR / "BAO_CAO_DO_AN_CHUONG_1_2_3_SEEE_HUST.docx"

CHAPTER_FILES = [
    OUT_DIR / "CHUONG_1_TONG_QUAN_HE_THONG.md",
    OUT_DIR / "CHUONG_2_PHAN_TICH_VA_THIET_KE_HE_THONG.md",
    OUT_DIR / "CHUONG_3_TRIEN_KHAI_VA_KET_QUA_GIAO_DIEN_MQTT.md",
]


def set_run_font(run, size_pt: float = 13.0, bold: bool = False, italic: bool = False, color: tuple[int, int, int] = (0, 0, 0)) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color)


def add_formatted_runs(p, text: str, base_size: float = 13.0, base_italic: bool = False) -> None:
    # Hỗ trợ in đậm **...**, in nghiêng *...*, mã `...`
    tokens = re.split(r"(\*\*.*?\*\*|\*[^\*]+?\*|`[^`]+?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2])
            set_run_font(r, size_pt=base_size, bold=True, italic=base_italic)
        elif tok.startswith("*") and tok.endswith("*"):
            r = p.add_run(tok[1:-1])
            set_run_font(r, size_pt=base_size, bold=False, italic=True)
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1])
            set_run_font(r, size_pt=base_size, bold=True, italic=base_italic)
        else:
            r = p.add_run(tok)
            set_run_font(r, size_pt=base_size, bold=False, italic=base_italic)


def build_cover_page(doc: Document) -> None:
    lines = [
        ("ĐẠI HỌC BÁCH KHOA HÀ NỘI", 14, True),
        ("TRƯỜNG ĐIỆN – ĐIỆN TỬ", 15, True),
        ("KHOA KỸ THUẬT ĐIỆN TỬ – VIỄN THÔNG", 13, True),
        ("--------------------  *  --------------------", 12, False),
        ("", 12, False),
        ("", 12, False),
        ("BÁO CÁO TIẾN ĐỘ ĐỒ ÁN TỐT NGHIỆP CỬ NHÂN", 18, True),
        ("(NỘI DUNG DUYỆT: CHƯƠNG 1, CHƯƠNG 2, CHƯƠNG 3)", 13, True),
        ("", 12, False),
        ("ĐỀ TÀI:", 14, True),
        (
            "NGHIÊN CỨU, THIẾT KẾ VÀ TRIỂN KHAI HỆ THỐNG NHÀ THÔNG MINH\n"
            "ĐIỀU KHIỂN BẰNG GIỌNG NÓI AI THỜI GIAN THỰC, PHÂN QUYỀN SINH TRẮC HỌC\n"
            "3 TẦNG, TỰ ĐỘNG HÓA 2 BỂ KẾT HỢP BẢN SAO SỐ 3D VÀ THU PHÁT\n"
            "SÓNG HỒNG NGOẠI TRÊN VI ĐIỀU KHIỂN ESP8266",
            14,
            True,
        ),
        ("", 12, False),
        ("", 12, False),
        ("Sinh viên thực hiện : NGUYỄN TUẤN ANH", 13, True),
        ("Ngành học           : Kỹ thuật Điện tử – Viễn thông", 13, False),
        ("Đơn vị              : Trường Điện – Điện tử (SEEE - HUST)", 13, False),
        ("", 12, False),
        ("", 12, False),
        ("HÀ NỘI – 2026", 13, True),
    ]
    for txt, sz, is_bold in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(txt)
        set_run_font(r, size_pt=sz, bold=is_bold)
    doc.add_page_break()


def render_table(doc: Document, table_lines: list[str]) -> None:
    rows_data = []
    for ln in table_lines:
        stripped = ln.strip().strip("|")
        cells = [c.strip().replace("<br>", "\n") for c in stripped.split("|")]
        if all(re.match(r"^[:\-]+$", c.replace(" ", "")) for c in cells if c):
            continue
        rows_data.append(cells)
    if not rows_data:
        return
    num_cols = max(len(r) for r in rows_data)
    tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows_data):
        for c_idx in range(num_cols):
            val = row[c_idx] if c_idx < len(row) else ""
            cell = tbl.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            if r_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                clean_v = re.sub(r"\*\*|\*|`|\\text\{([^}]*)\}|\$", r"\1", val)
                run = p.add_run(clean_v)
                set_run_font(run, size_pt=11.5, bold=True)
            else:
                clean_v = re.sub(r"\\pm", "±", val)
                clean_v = re.sub(r"\$", "", clean_v)
                add_formatted_runs(p, clean_v, base_size=11.5)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)


def parse_and_append_markdown(doc: Document, md_path: pathlib.Path, is_first_chapter: bool = False) -> None:
    if not is_first_chapter:
        doc.add_page_break()
    lines = md_path.read_text(encoding="utf-8").splitlines()
    in_code_block = False
    code_lang = ""
    code_buffer: list[str] = []
    table_buffer: list[str] = []

    for line in lines:
        s = line.strip()

        # Xử lý kết thúc bảng nếu dòng hiện tại không phải bảng
        if table_buffer and not s.startswith("|"):
            render_table(doc, table_buffer)
            table_buffer.clear()

        if s.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = s[3:].strip().lower()
                code_buffer.clear()
            else:
                in_code_block = False
                if code_lang == "mermaid":
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.left_indent = Cm(0.8)
                    p.paragraph_format.space_after = Pt(6)
                    r = p.add_run("[Sơ đồ khối / Lưu đồ thuật toán chi tiết: Xem trực quan trong bản Markdown hoặc hình vẽ tương ứng]\n" + "\n".join(code_buffer))
                    set_run_font(r, size_pt=10.5, italic=True, color=(40, 40, 40))
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.8)
                    p.paragraph_format.space_after = Pt(6)
                    r = p.add_run("\n".join(code_buffer))
                    set_run_font(r, size_pt=10.5)
            continue

        if in_code_block:
            code_buffer.append(line)
            continue

        if s.startswith("|"):
            table_buffer.append(s)
            continue

        if not s or s == "---":
            continue

        # Hình ảnh ![alt](path)
        m_img = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if m_img:
            img_rel = m_img.group(2).strip()
            img_path = (OUT_DIR / img_rel).resolve()
            if img_path.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(6)
                p_img.paragraph_format.space_after = Pt(3)
                run_img = p_img.add_run()
                run_img.add_picture(str(img_path), width=Cm(14.0))
            continue

        # Tiêu đề Chương (# CHƯƠNG ...)
        if s.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(14)
            r = p.add_run(s[2:].strip().upper())
            set_run_font(r, size_pt=16.0, bold=True)
            continue

        # Mục cấp 1 (## 1.1. ...)
        if s.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(s[3:].strip())
            set_run_font(r, size_pt=14.0, bold=True)
            continue

        # Mục cấp 2 (### 1.3.1. ...)
        if s.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(s[4:].strip().replace("`", ""))
            set_run_font(r, size_pt=13.0, bold=True)
            continue

        # Mục cấp 3 (#### a) ...)
        if s.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(s[5:].strip().replace("`", ""))
            set_run_font(r, size_pt=13.0, bold=True, italic=True)
            continue

        # Chú thích hình (*Hình 2.1. ...*)
        if s.startswith("*Hình ") and s.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run(s.strip("*"))
            set_run_font(r, size_pt=12.0, italic=True)
            continue

        # Tên bảng (**Bảng 1.1. ...**)
        if s.startswith("**Bảng ") and s.endswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(s.strip("*"))
            set_run_font(r, size_pt=12.0, bold=True)
            continue

        # Danh sách gạch đầu dòng hoặc đánh số
        clean_txt = (
            s.replace("$^\\circ\\text{C}$", "°C")
            .replace("$\\%\\text{RH}$", "%RH")
            .replace("$\\text{Hz}$", "Hz")
            .replace("$\\text{kHz}$", "kHz")
            .replace("$\\text{MHz}$", "MHz")
            .replace("$\\text{GHz}$", "GHz")
            .replace("$\\text{MB}$", "MB")
            .replace("$\\text{ms}$", "ms")
            .replace("$\\text{V DC}$", "V DC")
            .replace("$\\text{V}$", "V")
            .replace("$\\text{A}$", "A")
            .replace("$\\text{mA}$", "mA")
            .replace("$\\text{UTC}+7$", "UTC+7")
            .replace("$\\text{GMT}+7$", "GMT+7")
            .replace("$\\rightarrow$", "→")
            .replace("$\\leftrightarrow$", "↔")
            .replace("$\\pm 30$", "±30")
            .replace("$360^\\circ$", "360°")
            .replace("$90^\\circ$", "90°")
            .replace("$0^\\circ$", "0°")
            .replace("$F_0$", "F0")
            .replace("$L_2$", "L2")
            .replace("$$", "")
            .replace("$", "")
        )

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(6)
        if s.startswith(("- ", "+ ", "* ")) or re.match(r"^\d+\.\s", s):
            p.paragraph_format.left_indent = Cm(0.75)
        else:
            p.paragraph_format.first_line_indent = Cm(1.0)
        add_formatted_runs(p, clean_txt, base_size=13.0)

    if table_buffer:
        render_table(doc, table_buffer)


def main() -> None:
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)

    build_cover_page(doc)
    for idx, ch_file in enumerate(CHAPTER_FILES):
        parse_and_append_markdown(doc, ch_file, is_first_chapter=(idx == 0))

    doc.save(str(OUT_DOCX))
    print(f"Generated Word document successfully: {OUT_DOCX} ({OUT_DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
